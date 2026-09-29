#!/usr/bin/env python3
"""
clean_esi_schedule.py

Cleans an ESI "Schedule of Benefits Paid" .xls report (uneven merged cells,
repeated page headers, blank spacer rows, records split across 2-4 physical
rows) into a flat, filter/pivot-ready .xlsx file.

USAGE:
    python3 clean_esi_schedule.py input.xls output.xlsx

REQUIREMENTS:
    pip install xlrd openpyxl --break-system-packages   (if not already installed)

WHAT IT DOES:
    1. Reads the .xls with xlrd, including merged-cell ranges.
    2. Reconstructs a full grid where every cell in a merged range carries
       that merge's value (equivalent to "unmerge + fill").
    3. Auto-detects the repeated page-header blocks (wherever "Sl No." shows
       up as a literal label) and excludes them from the data.
    4. Auto-detects each beneficiary record's row-span from the merges in
       the "Sl No." column, and consolidates each record onto a single row,
       even where the original file scattered values across 2-4 rows or
       across duplicated/uneven merge widths for the same field.
    5. Builds flat column headers from the report's own multi-level header
       text, and drops columns that are pure formatting artifacts (no real
       sub-label and no data).
    6. Writes a new .xlsx with:
         - "Clean Data" sheet: one row per beneficiary, AutoFilter + frozen
           header, ready for PivotTable
         - "Summary (from report)" sheet: the report's own benefit-type
           totals table, preserved for cross-checking
         - "Notes" sheet: what was done + a validation check
    7. Validates itself: sums every "Amount" column in the cleaned data and
       compares it to the report's own printed Summary table. Prints a
       warning if anything doesn't match, instead of silently trusting it.

    Your original .xls file is never modified.
"""

import sys
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

try:
    import xlrd
except ImportError:
    sys.exit("Missing dependency. Run: pip install xlrd --break-system-packages")


def clean_text(s):
    """Collapse newlines/tabs/extra whitespace into single spaces."""
    if s is None:
        return ""
    return " ".join(str(s).replace("\n", " ").replace("\t", " ").split())


def to_number(v):
    """Convert numeric-looking strings to float; leave everything else as-is."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def build_filled_grid(sheet):
    """Build a full grid where every cell in a merged range holds that
    merge's value (mirrors what you'd see if you selected the merged
    range and hit 'Unmerge Cells' with the value copied into every cell)."""
    nrows, ncols = sheet.nrows, sheet.ncols
    grid = [
        [sheet.cell_value(r, c) if sheet.cell_value(r, c) != "" else None for c in range(ncols)]
        for r in range(nrows)
    ]
    for (r1, r2, c1, c2) in sheet.merged_cells:
        anchor = grid[r1][c1]
        for r in range(r1, r2):
            for c in range(c1, c2):
                grid[r][c] = anchor
    return grid


def find_header_markers(grid, marker_text="Sl No."):
    """Find the first row of every repeated page-header block. Because the
    label is inside a merge that was propagated down several rows, only the
    top row of each contiguous run of matches is a real block start."""
    nrows = len(grid)
    ncols = len(grid[0])
    raw_hits = []
    for r in range(nrows):
        for c in range(min(6, ncols)):
            if clean_text(grid[r][c]) == marker_text:
                raw_hits.append((r, c))
                break
    markers = [hit for i, hit in enumerate(raw_hits)
               if i == 0 or hit[0] != raw_hits[i - 1][0] + 1]
    return markers


def find_marker_row(grid, text, max_col=6):
    """Find the first row containing an exact-match label (e.g. 'Total :')."""
    for r in range(len(grid)):
        for c in range(min(max_col, len(grid[0]))):
            if clean_text(grid[r][c]).rstrip(":").strip().lower() == text.rstrip(":").strip().lower():
                return r
    return None


def header_for_column(grid, marker_row, c, id_col_ranges):
    """Build a flat header string for column c, using the report's own
    header text found relative to the nearest page-header marker row."""
    cat = clean_text(grid[marker_row][c])
    sub_item = clean_text(grid[marker_row + 1][c])   # e.g. 'Telephone Bill' under 'Miscellaneous'
    subtype = clean_text(grid[marker_row + 3][c])    # 'No. of Days' / 'Amount Rs.'

    for name, cols in id_col_ranges.items():
        if c in cols:
            return name

    if not cat:
        return None

    if cat.lower() == "miscellaneous" and sub_item:
        st = "No. of Days" if "day" in subtype.lower() else ("Amount Rs." if "amount" in subtype.lower() else "")
        return f"Misc - {sub_item} - {st}" if st else None

    if "day" in subtype.lower():
        return f"{cat} - No. of Days"
    if "amount" in subtype.lower():
        return f"{cat} - Amount Rs."
    return None  # bare category label with no real sub-column = formatting artifact, drop it


def clean_esi_file(input_path, output_path):
    wb_in = xlrd.open_workbook(input_path, formatting_info=True)
    sheet = wb_in.sheet_by_index(0)
    grid = build_filled_grid(sheet)
    ncols = len(grid[0])

    # ---- locate repeated page-header blocks ----
    markers = find_header_markers(grid)
    if not markers:
        sys.exit("Could not find 'Sl No.' header label - is this the same report template?")
    first_marker_row, sl_no_col = markers[0]

    # header block height = gap to the first actual data row after the marker
    first_data_row = None
    for (r1, r2, c1, c2) in sheet.merged_cells:
        if c1 <= sl_no_col < c2 and r1 > first_marker_row:
            first_data_row = r1
            break
    if first_data_row is None:
        sys.exit("Could not detect where beneficiary data begins.")
    header_height = first_data_row - first_marker_row

    excluded_rows = set()
    for (r, c) in markers:
        excluded_rows.update(range(r, r + header_height))

    # ---- locate end of data (the 'Total :' row) ----
    total_row = find_marker_row(grid, "Total :")
    if total_row is None:
        total_row = len(grid)

    # ---- identity columns (Sl No / Insurance No / Name / Daily Rate) ----
    id_col_ranges = {}
    id_labels = {
        "Sl No.": "sl no",
        "Insurance No.": "insurance no",
        "Name of Insured Person": "name of the insured",
        "Daily Benefit Rate (Rs.)": "daily benefit",
    }
    for c in range(1, ncols):
        cat = clean_text(grid[first_marker_row][c]).lower()
        for label, needle in id_labels.items():
            if needle in cat:
                id_col_ranges.setdefault(label, set()).add(c)

    # ---- record start rows (from merges in the Sl No. column) ----
    starts = set()
    for (r1, r2, c1, c2) in sheet.merged_cells:
        if c1 <= sl_no_col < c2 and r1 >= first_data_row and r1 < total_row:
            if r1 not in excluded_rows:
                starts.add(r1)
    starts = sorted(starts)
    if not starts:
        sys.exit("No beneficiary records detected.")

    # ---- build each record's row block, capped so it never swallows a header block ----
    blocks = []
    for i, s in enumerate(starts):
        nxt = starts[i + 1] if i + 1 < len(starts) else total_row
        end = nxt - 1
        for r in range(s + 1, end + 1):
            if r in excluded_rows:
                end = r - 1
                break
        blocks.append((s, end))

    # ---- build column -> header text map, then group columns by header ----
    col_header = {}
    for c in range(1, ncols):
        h = header_for_column(grid, first_marker_row, c, id_col_ranges)
        if h:
            col_header[c] = h

    header_cols = {}  # header text -> sorted list of column indices
    header_order = []
    for c in sorted(col_header):
        h = col_header[c]
        if h not in header_cols:
            header_cols[h] = []
            header_order.append(h)
        header_cols[h].append(c)

    # ---- extract one flat row per beneficiary ----
    records = []
    for (bs, be) in blocks:
        row_out = []
        for h in header_order:
            val = None
            for c in header_cols[h]:
                for r in range(bs, be + 1):
                    v = grid[r][c]
                    if v not in (None, ""):
                        val = v
                        break
                if val is not None:
                    break
            row_out.append(val)
        records.append(row_out)

    # ---- locate the report's own Summary table (for validation) ----
    summary_start = None
    for r in range(len(grid)):
        if any(clean_text(grid[r][c]).upper() == "SUMMARY" for c in range(min(6, ncols))):
            summary_start = r
            break
    summary_rows = []
    if summary_start is not None:
        # the summary table's label and number columns don't line up with
        # the main data table, so pull only columns that actually hold text
        # in this block: scan the header row (summary_start+1) for the
        # columns that carry real labels ('Nature of Benefit', 'No. of
        # Payments', etc.) and reuse exactly those columns for every row.
        label_row = summary_start + 1
        summary_cols = [c for c in range(ncols) if clean_text(grid[label_row][c])]
        # dedupe columns whose merge collapses to the same text as a
        # previous column (keep first occurrence of each distinct label)
        seen_labels = set()
        chosen_cols = []
        for c in summary_cols:
            lbl = clean_text(grid[label_row][c])
            if lbl not in seen_labels:
                seen_labels.add(lbl)
                chosen_cols.append(c)
        r = summary_start + 1
        prev_vals = None
        while r < len(grid):
            vals = [clean_text(grid[r][c]) for c in chosen_cols]
            if any(vals) and vals != prev_vals:
                summary_rows.append(vals)
                prev_vals = vals
            elif not any(vals) and summary_rows:
                break
            r += 1

    # ================= write output workbook =================
    out = Workbook()
    ws = out.active
    ws.title = "Clean Data"
    font_name = "Arial"

    hdr_font = Font(name=font_name, bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor="2F5597")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, h in enumerate(header_order, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = hdr_font, hdr_fill, border
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for i, rec in enumerate(records, start=2):
        for j, val in enumerate(rec, start=1):
            v = to_number(val)
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=font_name, size=10)
            c.border = border
            if j >= 4 and isinstance(v, float):
                c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="left" if j <= 3 else "right")

    ws.freeze_panes = "E2"
    ws.row_dimensions[1].height = 45
    widths = [6, 14, 24, 12] + [11] * (len(header_order) - 4)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.auto_filter.ref = ws.dimensions

    ws2 = out.create_sheet("Summary (from report)")
    for i, row in enumerate(summary_rows, start=1):
        for j, val in enumerate(row, start=1):
            c = ws2.cell(row=i, column=j, value=val)
            c.font = Font(name=font_name, size=10, bold=(i == 1 or i == len(summary_rows)))
            if i > 1 and j > 1:
                try:
                    c.value = float(val)
                    c.number_format = "#,##0.00"
                except ValueError:
                    pass
    for j, w in enumerate([32, 14, 16, 16], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w

    # ---- self-validation: compare extracted totals to the report's Summary ----
    amount_cols = [h for h in header_order if "Amount" in h]
    idx = {h: i for i, h in enumerate(header_order)}
    extracted_total = 0.0
    for h in amount_cols:
        for rec in records:
            v = to_number(rec[idx[h]])
            if isinstance(v, (int, float)):
                extracted_total += v

    report_total = None
    for row in summary_rows:
        if row[0].strip().upper().startswith("TOTAL"):
            for cell in reversed(row):
                try:
                    report_total = float(cell)
                    break
                except ValueError:
                    continue

    validation_note = "Not checked - no Summary table found in the source file."
    if report_total is not None:
        diff = round(extracted_total - report_total, 2)
        if abs(diff) < 0.01:
            validation_note = (
                f"PASSED - extracted total Rs. {extracted_total:,.2f} matches the "
                f"report's own printed total Rs. {report_total:,.2f} exactly."
            )
        else:
            validation_note = (
                f"MISMATCH - extracted total Rs. {extracted_total:,.2f} vs report's "
                f"printed total Rs. {report_total:,.2f} (difference: Rs. {diff:,.2f}). "
                f"Check this file's structure manually before trusting it."
            )

    ws3 = out.create_sheet("Notes")
    notes = [
        ["Source file", input_path],
        ["Records extracted", f"{len(records)} beneficiary payment records"],
        ["Method", "Merged cells were unmerged and their values propagated across the full "
                   "merge range; each beneficiary's data (originally split across multiple "
                   "physical rows) was consolidated onto a single flat row; repeated page "
                   "headers and blank spacer rows were removed."],
        ["Validation", validation_note],
        ["Original file", "Untouched - this is a separate cleaned copy."],
    ]
    for i, row in enumerate(notes, start=1):
        ws3.cell(row=i, column=1, value=row[0]).font = Font(name=font_name, bold=True, size=10)
        c2 = ws3.cell(row=i, column=2, value=row[1])
        c2.font = Font(name=font_name, size=10)
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[i].height = 30
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 90

    out.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Records extracted: {len(records)}")
    print(f"Validation: {validation_note}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 clean_esi_schedule.py input.xls output.xlsx")
        sys.exit(1)
    clean_esi_file(sys.argv[1], sys.argv[2])
  
